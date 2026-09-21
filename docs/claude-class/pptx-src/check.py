from pptx import Presentation
from pptx.util import Emu
import unicodedata, sys

p = Presentation('deck.pptx')
SW, SH = p.slide_width/914400, p.slide_height/914400
print(f"canvas {SW:.2f} x {SH:.2f} in, {len(p.slides)} slides\n")

def width_em(ch):
    # Korean/CJK ~1.0em, latin ~0.52em, space ~0.28em
    if ch == ' ': return 0.28
    return 1.0 if unicodedata.east_asian_width(ch) in ('W','F') else 0.52

problems = []
for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        if sh.left is None: continue
        x, y = sh.left/914400, sh.top/914400
        w, h = (sh.width or 0)/914400, (sh.height or 0)/914400
        if x < -0.01 or y < -0.01 or x+w > SW+0.01 or y+h > SH+0.01:
            problems.append(f"S{i}: OFF-CANVAS {sh.shape_type} at ({x:.2f},{y:.2f}) {w:.2f}x{h:.2f}")
        if not sh.has_text_frame: continue
        tf = sh.text_frame
        txt = tf.text
        if not txt.strip(): continue
        # font size: first run
        sz = None
        for para in tf.paragraphs:
            for r in para.runs:
                if r.font.size: sz = r.font.size.pt; break
            if sz: break
        if not sz: continue
        pad = 0.0 if tf.margin_left == 0 else 0.2
        avail_w = w - pad
        if avail_w <= 0: continue
        # estimate wrapped lines per hard-broken paragraph
        lines = 0
        for para_text in txt.split('\n'):
            if not para_text: lines += 1; continue
            em = sum(width_em(c) for c in para_text)
            line_w_em = avail_w * 72 / sz
            lines += max(1, -(-em // line_w_em))
        need = lines * sz * 1.45 / 72
        if need > h + 0.06:
            problems.append(f"S{i}: OVERFLOW? {sz}pt x{lines:.0f}ln needs {need:.2f}\" in {h:.2f}\" box "
                            f"({w:.2f}\" wide) :: {txt[:38]!r}")

if problems:
    print("\n".join(problems))
    print(f"\n{len(problems)} item(s) flagged")
else:
    print("no geometry or fit problems found")
